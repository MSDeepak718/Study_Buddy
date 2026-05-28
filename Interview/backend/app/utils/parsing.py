import logging
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)


def parse_pdf_pymupdf(file_path: str) -> str:
    """Parse PDF using PyMuPDF (fitz) - fast and reliable."""
    import fitz  # PyMuPDF

    text_parts = []
    try:
        doc = fitz.open(file_path)
        for page_num, page in enumerate(doc):
            text = page.get_text("text")
            if text.strip():
                text_parts.append(f"--- Page {page_num + 1} ---\n{text}")
        doc.close()
    except Exception as e:
        logger.error(f"PyMuPDF parsing failed for {file_path}: {e}")
        raise

    return "\n\n".join(text_parts)


def parse_pdf_pdfplumber(file_path: str) -> str:
    """Parse PDF using pdfplumber - better for table-heavy PDFs."""
    import pdfplumber

    text_parts = []
    try:
        with pdfplumber.open(file_path) as pdf:
            for page_num, page in enumerate(pdf.pages):
                text = page.extract_text()
                if text and text.strip():
                    text_parts.append(f"--- Page {page_num + 1} ---\n{text}")
    except Exception as e:
        logger.error(f"pdfplumber parsing failed for {file_path}: {e}")
        raise

    return "\n\n".join(text_parts)


def parse_docx(file_path: str) -> str:
    """Parse DOCX files using python-docx."""
    from docx import Document

    text_parts = []
    try:
        doc = Document(file_path)
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)

        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_parts.append(row_text)
    except Exception as e:
        logger.error(f"DOCX parsing failed for {file_path}: {e}")
        raise

    return "\n\n".join(text_parts)


def parse_document(file_path: str) -> str:
    """
    Parse a document based on its file extension.
    Uses fallback strategies for PDF parsing.

    Args:
        file_path: Path to the document file

    Returns:
        Extracted text content

    Raises:
        ValueError: If file type is not supported
    """
    path = Path(file_path)
    extension = path.suffix.lower()

    if extension == ".pdf":
        try:
            text = parse_pdf_pymupdf(file_path)
            if text.strip():
                logger.info(f"Successfully parsed PDF with PyMuPDF: {path.name}")
                return text
        except Exception:
            logger.warning(f"PyMuPDF failed, falling back to pdfplumber: {path.name}")

        # Fallback to pdfplumber
        text = parse_pdf_pdfplumber(file_path)
        logger.info(f"Successfully parsed PDF with pdfplumber: {path.name}")
        return text

    elif extension == ".docx":
        text = parse_docx(file_path)
        logger.info(f"Successfully parsed DOCX: {path.name}")
        return text

    elif extension == ".txt":
        with open(file_path, "r", encoding="utf-8") as f:
            text = f.read()
        logger.info(f"Successfully read TXT: {path.name}")
        return text

    else:
        raise ValueError(f"Unsupported file type: {extension}. Supported: .pdf, .docx, .txt")


def get_supported_extensions() -> list[str]:
    """Return list of supported file extensions."""
    return [".pdf", ".docx", ".txt"]
